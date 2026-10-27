"""
Turn the field team's spreadsheet into the shop file the app ships — and say exactly
what happened to every row.

RUN IT, IN THIS ORDER:

    python scripts/import_shops.py            1. read the sheet, write a REPORT and a
                                                 CANDIDATE file into scripts/out/.
                                                 The app is not touched.
    (read scripts/out/import-report.md)       2. a person reads the report.
    python scripts/import_shops.py --apply    3. only then: copy the candidate into
                                                 src/lib/data/shops.json.

FIELD SURVEY FILES from the app (ayos-survey-*.xlsx) are merged into the sheet first:

    python scripts/import_shops.py --survey ayos-survey-2026-10-12_1530.xlsx [more files]

which writes a merged COPY of the sheet into scripts/out/ (changed cells yellow, new
shops green, each with a comment saying what it was) and a merge report. A person checks
that copy, it becomes the field sheet, and then steps 1-3 run as usual. See "survey
merge" below for the rules.

WHY TWO STEPS. The shop list is field work collected shop by shop, and the app ships
it to every rider on the next update. A script that silently drops, merges or rewrites
rows is a script nobody can trust, so this one changes NOTHING the app reads until
someone has read what it did.

THE RULES, each one written down so a bad row is explained rather than lost:

    - line breaks and stray spaces are removed from every cell
    - "None", "N/A", "-" and empty cells all become "no value" (null), never the text
    - service tags: only dealer / repair / parts / accessories; blanks dropped
    - a row with no usable coordinates is SKIPPED and listed by row number
    - likely duplicates (within 75 m, similar name) are FLAGGED and BOTH are kept.
      Nothing is ever merged here: deciding two rows are one shop is a judgement
      about the real world, and the person who walked the street makes it.
    - phone numbers that do not parse are left out and listed for fixing
    - shop email addresses are NOT imported: the app has no use for them, and most are
      owners' personal addresses
    - photos are left empty
    - each shop's CITY is read from its address (one of Metro Manila's 17); a shop
      whose address names no city, or more than one, is listed for checking

Replaces scripts/convert_shops.py, which dropped duplicate rows without saying which.
"""

from __future__ import annotations

import argparse
import datetime as dt
import difflib
import hashlib
import json
import math
import os
import re
import shutil
import sys
from collections import Counter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

"""
The field spreadsheet lives OUTSIDE the repository, in a folder next to it, because it is
the field team's working file, not app code. Override with --xlsx for any other location.
"""
DEFAULT_XLSX = os.path.join(ROOT, os.pardir, "moto-app dataset", "Motorcycle Repairs_coordinates.xlsx")
DEFAULT_SHEET = "NCR"
OUT_DIR = os.path.join(HERE, "out")
CANDIDATE = os.path.join(OUT_DIR, "shops.candidate.json")
REPORT = os.path.join(OUT_DIR, "import-report.md")
SURVEY_REPORT = os.path.join(OUT_DIR, "survey-merge-report.md")
APP_FILE = os.path.join(ROOT, "src", "lib", "data", "shops.json")

"""
THE COLUMN MAP. Spreadsheet heading (lowercased, spaces squeezed) -> our field name.
When the sheet gains a column, this is the one place that changes.
"""
COLUMNS = {
    "name": "name",
    "address": "address",
    "latitude": "lat",
    "longitude": "lng",
    "services": "services",
    "contact no.": "phone",
    "social media": "social",
    "email": "email",
    "time": "time",
    "day closed": "closed",
    "image": "image",
    # Written by --survey; the import itself does not read it.
    "survey note": "note",
    "note": "note",
}

ALLOWED_TAGS = ("dealer", "repair", "parts", "accessories")

"""Metro Manila, drawn generously. Outside this is a typo, not a shop in another region."""
BOUNDS = {"lat": (14.30, 14.85), "lng": (120.85, 121.20)}

"""Spec section 5: warn when a shop sits within 75 m of another with a similar name."""
DUPLICATE_M = 75

"""What people type when a cell has nothing in it."""
EMPTY_WORDS = {"none", "n/a", "na", "null", "nil", "-", "--", "—", "?", "wala"}

DAYS = ["sunday", "monday", "tuesday", "wednesday", "thursday", "friday", "saturday"]

"""
METRO MANILA'S 17 CITIES (16 cities and the municipality of Pateros), as the app's ids
and the way they appear in addresses.

WHY FROM THE ADDRESS, NOT THE MAP POSITION: the app used to guess a shop's city from
which city CENTRE its pin was nearest to. Measured, that put 800 of 2,430
shops in a different city than their own address, and six cities had no group at all.
The address names exactly one city for 2,762 of 2,805 shops.

"Manila" must not match the "Metro Manila" that ends most addresses, hence the lookbehind.
"""
CITIES = [
    ("caloocan", r"caloocan"),
    ("las-pinas", r"las pi[nñ]as"),
    ("makati", r"makati"),
    ("malabon", r"malabon"),
    ("mandaluyong", r"mandaluyong"),
    ("manila", r"(?<!metro )(?<!metropolitan )\bmanila\b"),
    ("marikina", r"marikina"),
    ("muntinlupa", r"muntinlupa"),
    ("navotas", r"navotas"),
    ("paranaque", r"para[nñ]aque"),
    ("pasay", r"pasay"),
    ("pasig", r"pasig"),
    ("pateros", r"pateros"),
    ("quezon-city", r"quezon city|\bq\.? ?c\.?\b"),
    ("san-juan", r"san juan"),
    ("taguig", r"taguig"),
    ("valenzuela", r"valenzuela"),
]


# ------------------------------------------------------------------ cleaning a cell

def clean_text(value):
    """
    One cell, tidied. Line breaks and runs of spaces collapse to single spaces.
    Returns None for anything that means "nothing here", including the TEXT "None".
    """
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        # Excel stores a typed phone number as a float: 639281489926.0
        value = str(int(value))
    text = " ".join(str(value).split())
    if not text or text.lower().strip(" .,") in EMPTY_WORDS:
        return None
    return text


def to_float(value):
    try:
        return float(str(value).strip().rstrip(","))
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------ tags

def parse_tags(value):
    """
    "Parts, Accessories,, Repair"  ->  (["accessories", "parts", "repair"], [])

    Returns (tags, unknown words). Sorted, so two shops with the same services always
    produce the same list and a change in the output is a real change in the data.
    """
    text = clean_text(value) or ""
    tags, unknown = [], []
    for piece in re.split(r"[,/;]", text):
        word = piece.strip().lower()
        if not word:
            continue  # the blank between two commas
        if word in ALLOWED_TAGS:
            if word not in tags:
                tags.append(word)
        else:
            unknown.append(piece.strip())
    return sorted(tags), unknown


# ------------------------------------------------------------------ phones

def format_phone(raw):
    """
    One number in, one display string out, or None if it is not a number we can dial.

    Mobile:   0917-164-6472          (11 digits starting 09; +63 / 63 / 9... accepted)
    Landline: (02) 8537-6900         (Metro Manila numbers have been 8 digits since 2019)

    An extension ("loc. 1420") is cut off: a phone's dialler cannot press it for the
    rider, and the main line still reaches the shop.
    """
    text = re.split(r"\b(?:loc|local|ext)\b\.?", raw, flags=re.I)[0]
    digits = re.sub(r"\D", "", text)
    if digits.startswith("63"):
        digits = "0" + digits[2:]
    if len(digits) == 10 and digits.startswith("9"):
        digits = "0" + digits
    if len(digits) == 11 and digits.startswith("09"):
        return "%s-%s-%s" % (digits[:4], digits[4:7], digits[7:])
    if len(digits) == 10 and digits.startswith("02"):
        digits = digits[2:]
    if len(digits) == 9 and digits.startswith("2"):
        # "(02) 7001-8150" with its leading zero dropped, the way Excel drops it.
        digits = digits[1:]
    if len(digits) == 8 and digits[0] != "0":
        return "(02) %s-%s" % (digits[:4], digits[4:])
    return None


def parse_phones(value):
    """
    A contact cell can hold several numbers: "0917 1646472, 0916 492 8909".
    Returns (numbers we can dial, pieces we could not read), in the order written —
    the first one becomes the Call button.
    """
    text = clean_text(value)
    if not text:
        return [], []
    good, bad = [], []
    for piece in re.split(r"[,/;]|\bor\b", text):
        piece = piece.strip()
        if not re.search(r"\d", piece):
            continue
        number = format_phone(piece)
        if number is None:
            bad.append(piece)
        elif number not in good:
            good.append(number)
    return good, bad


# ------------------------------------------------------------------ links

def parse_links(value):
    """Web links only (Facebook, Shopee and so on). Trailing commas are stripped."""
    text = clean_text(value)
    if not text:
        return []
    links = []
    for token in re.split(r"[\s,]+", text):
        token = token.strip().rstrip(",;")
        if re.match(r"https?://\S+\.\S+", token) and token not in links:
            links.append(token)
    return links


# ------------------------------------------------------------------ hours

TIME_RANGE = re.compile(r"^(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})$")


def parse_time(value):
    """
    "9:00 - 20:00" -> ("09:00", "20:00").  "0:00 - 24:00" means open all day.
    Returns (open, close) or None when the cell is empty or unreadable.
    """
    text = clean_text(value)
    if not text:
        return None
    m = TIME_RANGE.match(text)
    if not m:
        return None
    oh, om, ch, cm = (int(x) for x in m.groups())
    if oh > 24 or ch > 24 or om > 59 or cm > 59:
        return None
    return "%02d:%02d" % (oh, om), "%02d:%02d" % (ch, cm)


def parse_closed(value):
    """
    The "Day Closed" column, as typed, into day numbers (0 = Sunday ... 6 = Saturday).

        "Everyday Open"          -> []
        "Sunday"                 -> [0]
        "Saturday and Sunday"    -> [0, 6]
        "Weekdays"               -> [1, 2, 3, 4, 5]
        "Only open on Sunday"    -> [1, 2, 3, 4, 5, 6]   (the opposite, said the other way)

    Returns None when the cell is empty or says something we cannot read, which the app
    treats as "we do not know", never as "open every day".
    """
    text = clean_text(value)
    if not text:
        return None
    low = text.lower()
    if low in ("everyday open", "open everyday", "open every day", "everyday"):
        return []
    if low == "weekdays":
        return [1, 2, 3, 4, 5]
    if low == "weekends":
        return [0, 6]
    only = re.match(r"^only open on (\w+)$", low)
    if only and only.group(1) in DAYS:
        keep = DAYS.index(only.group(1))
        return [d for d in range(7) if d != keep]
    words = [w for w in re.split(r"[\s,/&]+|\band\b", low) if w]
    if words and all(w in DAYS for w in words):
        return sorted({DAYS.index(w) for w in words})
    return None


# ------------------------------------------------------------------ city

STREET_NAME_CITIES = {"manila", "san-juan"}


def city_from_address(address):
    """
    ("quezon-city", "one")      the address names exactly one city
    ("manila", "several")       it names more than one; the LAST is taken, because a
                                Philippine address runs street -> barangay -> city, so
                                an earlier match is usually a street ("Quezon Ave")
    (None, "none")              it names none
    """
    text = (address or "").lower()
    hits = []
    for slug, pattern in CITIES:
        for m in re.finditer(pattern, text):
            hits.append((m.start(), slug))
    """
    "Manila" and "San Juan" are also STREET names all over Metro Manila (Manila South
    Road, Manila Gravel Pit Road, San Juan Evangelista Street), so when another city is
    named too, they are taken to be the street, not the city.
    """
    found = {slug for _, slug in hits}
    if len(found) > 1:
        kept = [(pos, slug) for pos, slug in hits if slug not in STREET_NAME_CITIES]
        if kept:
            hits = kept
            found = {slug for _, slug in hits}
    if not found:
        return None, "none"
    last = max(hits)[1]
    return last, ("one" if len(found) == 1 else "several")


def nearest_city(lat, lng, centres):
    """For an address with no city: the city whose shops' middle point is nearest."""
    return min(centres, key=lambda c: (centres[c][0] - lat) ** 2 + (centres[c][1] - lng) ** 2)


def city_centres(shops):
    """The middle (median) of each city's clearly-addressed shops, as (lat, lng)."""
    by_city = {}
    for s in shops:
        if s.get("_city_status") == "one":
            by_city.setdefault(s["city"], []).append((s["lat"], s["lng"]))
    centres = {}
    for city, pts in by_city.items():
        lats = sorted(p[0] for p in pts)
        lngs = sorted(p[1] for p in pts)
        centres[city] = (lats[len(lats) // 2], lngs[len(lngs) // 2])
    return centres


# ------------------------------------------------------------------ duplicates

GENERIC = re.compile(
    r"\b(motor ?(cycle)?s?|moto|shop|parts?|motorparts|accessor(ies|y)|accesories|and|"
    r"repair|services?|center|centre|trading|enterprises?|inc|corp|store|supply|supplies|"
    r"genuine|authorized|the|of|branch|main)\b"
)


def name_core(name):
    """
    The part of a shop name that tells it apart from its neighbours.

    Without this, "Pega Motorcycle Parts and Accessories" and "Marklhyne Motorcycle
    Parts and Accessories" score as near-identical, because most of each name is the
    same generic words. Measured on this dataset: comparing whole names raises 54
    flags; comparing only the distinctive part raises 18, all of them worth a look.
    """
    text = re.sub(r"[&.\-]", " ", name.lower())
    text = GENERIC.sub(" ", text)
    return re.sub(r"[^a-z0-9]+", "", text)


def similar_names(a, b):
    ca, cb = name_core(a), name_core(b)
    if not ca or not cb:
        return False
    if difflib.SequenceMatcher(None, ca, cb).ratio() >= 0.8:
        return True
    return min(len(ca), len(cb)) >= 4 and (ca in cb or cb in ca)


def metres(lat1, lng1, lat2, lng2):
    r = 6371000
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lng2 - lng1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(h))


def find_duplicates(shops):
    """
    Every pair within 75 m with a similar name, closest first. Sorted by latitude so
    each shop is only compared with the few that could possibly be that close.
    """
    ordered = sorted(shops, key=lambda s: s["lat"])
    window = DUPLICATE_M / 111_320
    pairs = []
    for i, a in enumerate(ordered):
        for b in ordered[i + 1:]:
            if b["lat"] - a["lat"] > window:
                break
            d = metres(a["lat"], a["lng"], b["lat"], b["lng"])
            if d <= DUPLICATE_M and similar_names(a["name"], b["name"]):
                pairs.append((round(d), a, b))
    pairs.sort(key=lambda p: (p[0], p[1]["_row"]))
    return pairs


# ------------------------------------------------------------------ ids

def make_id(name, lat, lng, taken):
    """
    Derived from the shop itself, not its row number.

    The old converter numbered shops by position, so inserting one row near the top of
    the sheet renumbered every shop below it. This id survives rows being added,
    removed or re-sorted. It changes only if the shop's name or position is corrected,
    which is a real change to that shop.
    """
    key = "%s|%.4f|%.4f" % (" ".join(name.lower().split()), lat, lng)
    base = "s" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:8]
    sid, n = base, 2
    while sid in taken:
        sid = "%s-%d" % (base, n)
        n += 1
    taken.add(sid)
    return sid


# ------------------------------------------------------------------ read + build

def read_sheet(path, sheet):
    import openpyxl  # imported here so the tests run without it

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    return read_rows(ws), ws.title


def read_rows(ws):
    """Every non-empty row of a worksheet as {field: value, "_row": sheet row number}."""
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    index = {}
    for pos, cell in enumerate(header):
        key = " ".join(str(cell or "").lower().split())
        if key in COLUMNS:
            index[COLUMNS[key]] = pos
    for need in ("name", "lat", "lng"):
        if need not in index:
            sys.exit("The sheet has no '%s' column." % need)
    out = []
    for n, row in enumerate(rows, start=2):
        if not row or not any(c not in (None, "") for c in row):
            continue
        rec = {f: (row[p] if p < len(row) else None) for f, p in index.items()}
        rec["_row"] = n
        out.append(rec)
    return out


def build(rows):
    """Every row in, (shops, report facts) out. Pure: no files, so it can be tested."""
    shops, skipped, bad_phones, unknown_tags = [], [], [], Counter()
    no_tags, odd_closed, hours_mismatch = [], [], []
    counts = Counter()
    taken = set()

    for row in rows:
        n = row["_row"]
        name = clean_text(row.get("name"))
        if not name:
            skipped.append((n, "no name", ""))
            continue
        lat, lng = to_float(row.get("lat")), to_float(row.get("lng"))
        if lat is None or lng is None:
            skipped.append((n, "no coordinates", name))
            continue
        if not (BOUNDS["lat"][0] <= lat <= BOUNDS["lat"][1] and BOUNDS["lng"][0] <= lng <= BOUNDS["lng"][1]):
            skipped.append((n, "outside Metro Manila", name))
            continue

        tags, unknown = parse_tags(row.get("services"))
        unknown_tags.update(unknown)
        if not tags:
            no_tags.append((n, name))

        shop = {
            "_row": n,
            "id": make_id(name, lat, lng, taken),
            "name": name,
            "lat": round(lat, 6),
            "lng": round(lng, 6),
            "tags": tags,
        }
        address = clean_text(row.get("address"))
        if address:
            shop["address"] = address
        city, status = city_from_address(address)
        shop["city"] = city
        shop["_city_status"] = status

        phones, unreadable = parse_phones(row.get("phone"))
        if phones:
            shop["phones"] = phones
        for piece in unreadable:
            bad_phones.append((n, name, piece))

        links = parse_links(row.get("social"))
        if links:
            shop["links"] = links

        time = parse_time(row.get("time"))
        closed = parse_closed(row.get("closed"))
        if clean_text(row.get("closed")) and closed is None:
            odd_closed.append((n, name, clean_text(row.get("closed"))))
        if (time is None) != (clean_text(row.get("closed")) is None):
            hours_mismatch.append((n, name))
        if time or closed is not None:
            hours = {}
            if time:
                hours["open"], hours["close"] = time
            if closed is not None:
                hours["closed"] = closed
            shop["hours"] = hours

        if clean_text(row.get("email")):
            counts["email present, not imported"] += 1
        if clean_text(row.get("image")):
            counts["image present, not imported"] += 1

        shops.append(shop)

    """
    Second pass for addresses that named no city: the nearest city by where the clearly
    addressed shops actually are. Both these and the "several" ones are listed in the
    report so the sheet can be corrected.
    """
    centres = city_centres(shops)
    city_unclear = []
    for s in shops:
        if s["_city_status"] == "none":
            s["city"] = nearest_city(s["lat"], s["lng"], centres) if centres else None
        if s["_city_status"] != "one":
            city_unclear.append((s["_row"], s["name"], s.get("address", ""), s["city"],
                                 "no city in address" if s["_city_status"] == "none" else "names more than one city"))

    facts = {
        "city_unclear": city_unclear,
        "city_centres": centres,
        "skipped": skipped,
        "bad_phones": bad_phones,
        "unknown_tags": unknown_tags,
        "no_tags": no_tags,
        "odd_closed": odd_closed,
        "hours_mismatch": hours_mismatch,
        "duplicates": find_duplicates(shops),
        "generic_names": [(s["_row"], s["name"]) for s in shops if not name_core(s["name"])],
        "counts": counts,
    }
    return shops, facts


def compare_with_app(shops, app_file):
    """Shops in the file the app ships today that the new sheet no longer has, and vice versa."""
    if not os.path.exists(app_file):
        return None
    with open(app_file, encoding="utf-8") as fh:
        current = json.load(fh)
    current = current["shops"] if isinstance(current, dict) else current
    key = lambda s: (s["name"].strip().lower(), round(s["lat"], 4), round(s["lng"], 4))
    now = Counter(key(s) for s in current)
    new = Counter(key(s) for s in shops)
    missing = now - new
    """
    One row per MISSING COPY. When the sheet had a shop twice and one copy was deleted, only
    that one copy is gone; listing every row with the same name and position reported both
    copies as disappearing (seen 11 Oct 2026: one KRD duplicate shown as two shops gone).
    """
    gone = []
    left = dict(missing)
    for s in current:
        k = key(s)
        if left.get(k, 0) > 0:
            gone.append((s["name"], round(s["lat"], 4), round(s["lng"], 4)))
            left[k] -= 1
    gone.sort()
    added = sum((new - now).values())
    return {"current": len(current), "gone": gone, "added": added, "kept": sum((now & new).values())}


# ------------------------------------------------------------------ report

def write_report(path, source, sheet, collected_at, rows, shops, facts, diff, size):
    def table(head, body):
        lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
        lines += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in body]
        return "\n".join(lines)

    with_phone = sum(1 for s in shops if "phones" in s)
    with_hours = sum(1 for s in shops if "hours" in s)
    with_links = sum(1 for s in shops if "links" in s)
    L = []
    L.append("# Shop import report")
    L.append("")
    L.append("Generated %s from `%s`, sheet `%s`." % (dt.date.today().isoformat(), os.path.basename(source), sheet))
    L.append("List date (shown to riders as \"list from\"): **%s**." % collected_at)
    L.append("")
    L.append("**Nothing in the app has changed yet.** Read this, then run `python scripts/import_shops.py --apply`.")
    L.append("")
    L.append("## Totals")
    L.append("")
    L.append(table(["", "count"], [
        ["rows in the sheet", len(rows)],
        ["**imported**", "**%d**" % len(shops)],
        ["skipped", len(facts["skipped"])],
        ["flagged as possible duplicates (pairs, both kept)", len(facts["duplicates"])],
        ["with at least one phone number", with_phone],
        ["with opening hours or closed days", with_hours],
        ["with a Facebook or other web link", with_links],
        ["shop emails left out on purpose", facts["counts"]["email present, not imported"]],
        ["city unclear in the address (listed below)", len(facts["city_unclear"])],
        ["photos", 0],
        ["candidate file size", "%d KB" % round(size / 1024)],
    ]))
    L.append("")
    if diff:
        L.append("## Compared with the list the app ships today")
        L.append("")
        L.append("%d shops today. %d stay, %d are new, **%d are no longer in the sheet** and will disappear from the app:" %
                 (diff["current"], diff["kept"], diff["added"], len(diff["gone"])))
        L.append("")
        L.append(table(["name", "lat", "lng"], diff["gone"]) if diff["gone"] else "_none_")
        L.append("")
        L.append("Some of these may be renames or moved pins rather than closures. If any should stay, fix the sheet and run the script again.")
        L.append("")
    L.append("## Skipped (not imported)")
    L.append("")
    L.append(table(["sheet row", "reason", "name"], facts["skipped"]) if facts["skipped"] else "_none_")
    L.append("")
    L.append("## Possible duplicates: BOTH are imported, nothing was merged")
    L.append("")
    L.append("Within %d m of each other with a similar name. Check each on the ground or on the map; if two rows are one shop, delete one row in the sheet and run again." % DUPLICATE_M)
    L.append("")
    L.append(table(["metres", "row", "name", "row", "name"],
                   [(d, a["_row"], a["name"], b["_row"], b["name"]) for d, a, b in facts["duplicates"]])
             if facts["duplicates"] else "_none_")
    L.append("")
    if facts["generic_names"]:
        L.append("Named only with generic words, so they cannot be checked for duplicates by name:")
        L.append("")
        L.append(table(["row", "name"], facts["generic_names"]))
        L.append("")
    L.append("## Phone numbers that could not be read (left out)")
    L.append("")
    L.append("A 7-digit Metro Manila landline stopped working on 6 Oct 2019, when every 02 number gained a "
             "carrier digit in front (8 for PLDT, 7 for Globe, 3 for Bayan). We cannot tell which one a shop "
             "uses, so these need a quick call or visit to get the current number.")
    L.append("")
    L.append(table(["row", "shop", "as typed"], facts["bad_phones"]) if facts["bad_phones"] else "_none_")
    L.append("")
    L.append("## City unclear in the address")
    L.append("")
    L.append("The app files these under the city shown. Fix the address in the sheet if it is wrong.")
    L.append("")
    L.append(table(["row", "shop", "address", "filed under", "why"], facts["city_unclear"])
             if facts["city_unclear"] else "_none_")
    L.append("")
    L.append("## Other things to fix in the sheet")
    L.append("")
    L.append("**Rows with no service tag** (imported with no tags):")
    L.append("")
    L.append(table(["row", "name"], facts["no_tags"]) if facts["no_tags"] else "_none_")
    L.append("")
    if facts["unknown_tags"]:
        L.append("**Service words that are not dealer / repair / parts / accessories** (dropped):")
        L.append("")
        L.append(table(["word", "times"], facts["unknown_tags"].most_common()))
        L.append("")
    L.append("**\"Day Closed\" values we could not read** (stored as unknown):")
    L.append("")
    L.append(table(["row", "name", "as typed"], facts["odd_closed"]) if facts["odd_closed"] else "_none_")
    L.append("")
    L.append("**Hours filled in but Day Closed empty, or the other way round:**")
    L.append("")
    L.append(table(["row", "name"], facts["hours_mismatch"]) if facts["hours_mismatch"] else "_none_")
    L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


# ------------------------------------------------------------------ survey merge

"""
MERGING A FIELD SURVEY INTO THE SHEET (--survey).

The app's survey export is one .xlsx with a "Changes" tab (every edit: shop id, column,
before, after) and a "New shops" tab (rows in the sheet's columns). This folds both into
a COPY of the field sheet, so nobody retypes anything:

    - an edit is applied only while the sheet still says what the surveyor saw
      ("before"). If someone changed that cell since, it is a CONFLICT: left alone and
      listed, never decided by the script;
    - values are compared the way the app reads them, so "0915 9777778" and
      "0915-977-7778" are the same number and only real changes count;
    - email and note were never in the app, so the surveyor had no "before" for them:
      they fill an empty cell, and anything already there is a conflict;
    - a shop whose id is not in the sheet any more (renamed, moved or deleted since) is
      listed, not guessed at;
    - new shops are appended at the bottom; one already in the sheet (same name, same
      spot) is not added twice, so merging the same file again is harmless;
    - every cell the merge writes is coloured (yellow changed, green added) and carries
      a comment saying what it was before, so the person checking sees exactly what moved.

Nothing the app reads is touched. The merged copy is checked by a person, becomes the
field sheet, and only then goes through the normal import and --apply.

plan_survey() decides everything on plain rows, so it is tested without a spreadsheet
library; write_survey() only copies its decisions into the workbook.
"""

SURVEY_COLUMNS = {
    "name": "name",
    "address": "address",
    "services": "services",
    "contact no.": "phone",
    "social media": "social",
    "email": "email",
    "time": "time",
    "day closed": "closed",
    "note": "note",
    "latitude, longitude": "location",
}
"""Fields the app never had, so the survey's "before" for them is always blank."""
UNKNOWN_TO_APP = ("email", "note")
FIELD_HEADING = {"name": "Name", "address": "Address", "lat": "Latitude", "lng": "Longitude",
                 "services": "Services", "phone": "Contact No.", "social": "Social Media",
                 "email": "Email", "time": "Time", "closed": "Day Closed", "note": "Survey note"}
"""About 1 m of latitude. Closer than that, two coordinates are the same pin."""
SAME_SPOT = 1e-5
FILL_CHANGED = "FFFFE58A"
FILL_ADDED = "FFC9EBD3"


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_location(text):
    """'14.6531200, 121.0345600' -> (lat, lng), or None."""
    parts = [p.strip() for p in str(text or "").split(",")]
    if len(parts) != 2:
        return None
    lat, lng = to_float(parts[0]), to_float(parts[1])
    return None if lat is None or lng is None else (lat, lng)


def same_value(field, a, b):
    """True when two cells say the same thing as far as the app is concerned."""
    if field == "phone":
        return parse_phones(a)[0] == parse_phones(b)[0]
    if field == "services":
        return parse_tags(a)[0] == parse_tags(b)[0]
    if field == "social":
        return parse_links(a) == parse_links(b)
    if field == "time":
        return parse_time(a) == parse_time(b)
    if field == "closed":
        return parse_closed(a) == parse_closed(b)
    norm = lambda v: " ".join((clean_text(v) or "").lower().split())
    return norm(a) == norm(b)


def read_survey(path):
    """(changes, new shop records) from one app survey export."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    if "Changes" not in wb.sheetnames or "New shops" not in wb.sheetnames:
        sys.exit("%s is not an Ayos survey export (it has no Changes and New shops tabs)." % path)

    def records(ws):
        rows = ws.iter_rows(values_only=True)
        head = [" ".join(str(c or "").lower().split()) for c in next(rows, ())]
        return [dict(zip(head, r)) for r in rows if r and any(c not in (None, "") for c in r)]

    changes = []
    for r in records(wb["Changes"]):
        field = SURVEY_COLUMNS.get(" ".join(cell_text(r.get("column")).lower().split()))
        if field:
            changes.append({
                "date": cell_text(r.get("date")),
                "shop_id": cell_text(r.get("shop id")),
                "shop": cell_text(r.get("shop name")),
                "column": cell_text(r.get("column")),
                "field": field,
                "before": cell_text(r.get("before")),
                "after": cell_text(r.get("after")),
            })
    added = []
    for r in records(wb["New shops"]):
        rec = {COLUMNS[k]: v for k, v in r.items() if k in COLUMNS}
        rec["note"] = r.get("note")
        rec["added_on"] = cell_text(r.get("added on"))
        added.append(rec)
    return changes, added


def collapse(changes):
    """One change per shop and column: the FIRST before and the LAST after."""
    out = {}
    for c in sorted(changes, key=lambda c: c["date"]):  # stable: same-minute edits keep file order
        key = (c["shop_id"], c["field"])
        if key in out:
            out[key]["after"], out[key]["date"] = c["after"], c["date"]
        else:
            out[key] = dict(c)
    return list(out.values())


def plan_survey(rows, surveys):
    """
    Decide every write. `rows` are read_rows() output; `surveys` is [(source name,
    changes, new shops)]. Returns (writes, appends, facts, merged rows), where a write is
    (sheet row, field, new value, comment) and an append is (sheet row, values by field,
    comment). New rows go straight after the last row that has anything in it.
    """
    rows = [dict(r) for r in rows]
    by_row = {r["_row"]: r for r in rows}
    shops, _ = build(rows)
    row_of = {s["id"]: s["_row"] for s in shops}
    """
    A shop's id comes from its name and position, so a pin moved by an earlier merge gives
    it a new id. Fallback: the one row with the same name. Only when the name is unique,
    and the before-check below still guards every write.
    """
    by_name = {}
    for s in shops:
        by_name.setdefault(" ".join(s["name"].lower().split()), []).append(s["_row"])
    next_row = max([r["_row"] for r in rows] + [1]) + 1

    writes, appends = [], []
    facts = {"applied": [], "already": [], "conflicts": [], "missing": [], "unreadable": [],
             "added": [], "already_added": [], "by_name": []}

    for source, changes, added in surveys:
        for c in collapse(changes):
            n = row_of.get(c["shop_id"])
            if n is None:
                twins = by_name.get(" ".join(c["shop"].lower().split()), [])
                if len(twins) == 1:
                    n = twins[0]
                    if not any(r == n for r, _, _ in facts["by_name"]):
                        facts["by_name"].append((n, c["shop"], c["shop_id"]))
            if n is None:
                facts["missing"].append((c["shop"], c["shop_id"], c["column"], c["after"] or "(blank)", source))
                continue
            row = by_row[n]
            name = clean_text(row.get("name")) or c["shop"]
            note = "Survey %s (%s): was " % (c["date"], source)
            shown = lambda v: clean_text(v) or "(blank)"

            if c["field"] == "location":
                after = parse_location(c["after"])
                before = parse_location(c["before"])
                now = (to_float(row.get("lat")), to_float(row.get("lng")))
                if after is None:
                    facts["unreadable"].append((n, name, c["column"], c["after"]))
                    continue
                close = lambda p, q: None not in (p[0], q[0]) and abs(p[0] - q[0]) <= SAME_SPOT and abs(p[1] - q[1]) <= SAME_SPOT
                if close(now, after):
                    facts["already"].append((n, name, c["column"]))
                    continue
                if before is not None and not close(now, before):
                    facts["conflicts"].append((n, name, c["column"], "%s, %s" % now, c["before"], c["after"], source))
                    continue
                for field, value, was in (("lat", after[0], now[0]), ("lng", after[1], now[1])):
                    writes.append((n, field, value, note + str(was)))
                    row[field] = value
                facts["applied"].append((n, name, c["column"], "%s, %s" % now, c["after"]))
                continue

            field = c["field"]
            now = row.get(field)
            if same_value(field, now, c["after"]) and cell_text(now):
                facts["already"].append((n, name, c["column"]))
                continue
            if field in UNKNOWN_TO_APP:
                fits = clean_text(now) is None
            else:
                fits = same_value(field, now, c["before"])
            if not fits:
                facts["conflicts"].append((n, name, c["column"], shown(now),
                                           c["before"] or "(blank)", c["after"] or "(blank)", source))
                continue
            value = c["after"] or None
            writes.append((n, field, value, note + shown(now)))
            row[field] = value
            facts["applied"].append((n, name, c["column"], shown(now), c["after"] or "(blank)"))

        for rec in added:
            name = clean_text(rec.get("name"))
            lat, lng = to_float(rec.get("lat")), to_float(rec.get("lng"))
            if not name or lat is None or lng is None:
                facts["unreadable"].append(("new", name or "(no name)", "New shop", "missing name or coordinates"))
                continue
            twin = next((r for r in rows if same_value("name", r.get("name"), name)
                         and to_float(r.get("lat")) is not None
                         and abs(to_float(r.get("lat")) - lat) <= SAME_SPOT
                         and abs(to_float(r.get("lng")) - lng) <= SAME_SPOT), None)
            if twin:
                facts["already_added"].append((twin["_row"], name))
                continue
            values = {f: rec.get(f) for f in ("name", "address", "lat", "lng", "services", "phone",
                                               "social", "email", "time", "closed", "note")}
            values["lat"], values["lng"] = lat, lng
            new_row = dict(values, _row=next_row)
            rows.append(new_row)
            appends.append((next_row, values, "Added by survey %s (%s)" % (rec.get("added_on") or "", source)))
            facts["added"].append((next_row, name))
            next_row += 1

    return writes, appends, facts, rows


def write_survey(ws, writes, appends):
    """Copy plan_survey's decisions into the worksheet: values, colours and comments."""
    from openpyxl.comments import Comment
    from openpyxl.styles import PatternFill

    col = {}
    last = 0
    for i, cell in enumerate(ws[1], start=1):
        key = " ".join(str(cell.value or "").lower().split())
        if key:
            last = i
        if key in COLUMNS and COLUMNS[key] not in col:
            col[COLUMNS[key]] = i

    def column(field):
        nonlocal last
        if field not in col:  # only the survey note column can be missing
            last += 1
            ws.cell(1, last, FIELD_HEADING[field])
            col[field] = last
        return col[field]

    changed = PatternFill("solid", fgColor=FILL_CHANGED)
    new = PatternFill("solid", fgColor=FILL_ADDED)
    for n, field, value, note in writes:
        cell = ws.cell(n, column(field))
        cell.value = value
        cell.fill = changed
        cell.comment = Comment(note, "Ayos survey")
    for row, values, note in appends:
        for field, value in values.items():
            if value in (None, ""):
                continue
            cell = ws.cell(row, column(field))
            cell.value = value
            cell.fill = new
        ws.cell(row, column("name")).comment = Comment(note, "Ayos survey")


def write_survey_report(path, master, sources, facts, checks):
    def table(head, body):
        lines = ["| " + " | ".join(head) + " |", "|" + "---|" * len(head)]
        lines += ["| " + " | ".join(str(c).replace("|", "/") for c in r) + " |" for r in body]
        return "\n".join(lines)

    L = ["# Survey merge report", ""]
    L.append("Generated %s. Field sheet `%s`, survey files: %s." % (
        dt.date.today().isoformat(), os.path.basename(master), ", ".join("`%s`" % s for s in sources)))
    L.append("")
    L.append("**Nothing in the app has changed.** The merged sheet is a copy: check the coloured cells, "
             "fix anything wrong, make it the field sheet, then run the import as usual.")
    L.append("")
    L.append(table(["", "count"], [
        ["changes applied (yellow)", len(facts["applied"])],
        ["new shops added (green, at the bottom)", len(facts["added"])],
        ["**conflicts, left for a person**", "**%d**" % len(facts["conflicts"])],
        ["shops not found in the sheet", len(facts["missing"])],
        ["already in the sheet, nothing to do", len(facts["already"]) + len(facts["already_added"])],
        ["unreadable", len(facts["unreadable"])],
    ]))
    L.append("")
    sections = [
        ("Conflicts: the sheet changed since the surveyor saw it",
         "Nothing was written. Decide which value is right and fix the cell by hand.",
         ["row", "shop", "column", "sheet says now", "surveyor saw", "surveyor wrote", "file"], facts["conflicts"]),
        ("Shops not found in the sheet",
         "Renamed, moved or deleted since the app's list was made. Find them by hand.",
         ["shop", "id", "column", "surveyor wrote", "file"], facts["missing"]),
        ("Matched by name, not id",
         "The id was not in the sheet (usually a pin moved earlier), but exactly one row has this name. Check it is the same shop.",
         ["row", "shop", "id in the survey"], facts["by_name"]),
        ("Applied", "", ["row", "shop", "column", "was", "now"], facts["applied"]),
        ("New shops added", "", ["row", "name"], facts["added"]),
        ("Unreadable", "", ["row", "shop", "column", "value"], facts["unreadable"]),
    ]
    for title, note, head, body in sections:
        L.append("## " + title)
        L.append("")
        if note:
            L.append(note)
            L.append("")
        L.append(table(head, body) if body else "_none_")
        L.append("")
    L.append("## Problems in the rows the survey touched")
    L.append("")
    L.append("The same checks the import runs, limited to changed and added rows.")
    L.append("")
    L.append(table(["kind", "row", "detail"], checks) if checks else "_none_")
    L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def survey_checks(rows, touched):
    """The import's own checks, kept to the rows the survey wrote."""
    _, facts = build(rows)
    out = []
    out += [("skipped: " + why, n, name) for n, why, name in facts["skipped"] if n in touched]
    out += [("phone not readable", n, "%s: %s" % (name, piece)) for n, name, piece in facts["bad_phones"] if n in touched]
    out += [("possible duplicate", a["_row"], "%s / row %d %s (%d m)" % (a["name"], b["_row"], b["name"], d))
            for d, a, b in facts["duplicates"] if a["_row"] in touched or b["_row"] in touched]
    out += [("city unclear", n, "%s: %s" % (name, why)) for n, name, _, _, why in facts["city_unclear"] if n in touched]
    out += [("Day Closed not readable", n, "%s: %s" % (name, v)) for n, name, v in facts["odd_closed"] if n in touched]
    return out


def merge_survey(master, sheet, paths):
    import openpyxl

    wb = openpyxl.load_workbook(master)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    if any(c.data_type == "f" for r in ws.iter_rows() for c in r):
        sys.exit("The sheet has formulas; merging would read them as text. Paste them as values first.")
    rows = read_rows(ws)
    surveys = [(os.path.basename(p),) + read_survey(p) for p in paths]
    writes, appends, facts, merged = plan_survey(rows, surveys)
    write_survey(ws, writes, appends)

    os.makedirs(OUT_DIR, exist_ok=True)
    stem = os.path.splitext(os.path.basename(master))[0]
    out = os.path.join(OUT_DIR, "%s (survey merged %s).xlsx" % (stem, dt.date.today().isoformat()))
    wb.save(out)
    touched = {n for n, *_ in writes} | {n for n, _ in facts["added"]}
    write_survey_report(SURVEY_REPORT, master, [s[0] for s in surveys], facts, survey_checks(merged, touched))

    print("%d applied, %d new shops, %d conflicts, %d not found, %d already in the sheet" % (
        len(facts["applied"]), len(facts["added"]), len(facts["conflicts"]), len(facts["missing"]),
        len(facts["already"]) + len(facts["already_added"])))
    print("merged sheet  %s" % out)
    print("report        %s" % SURVEY_REPORT)
    print("Nothing in the app changed. Check the coloured cells, make the merged copy the field sheet,")
    print("then run the import as usual.")


# ------------------------------------------------------------------ main

def main():
    ap = argparse.ArgumentParser(description="Spreadsheet of shops -> report + candidate app file.")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--sheet", default=DEFAULT_SHEET)
    ap.add_argument("--collected-at", help="YYYY-MM-DD. Defaults to the day the spreadsheet was last saved.")
    ap.add_argument("--apply", action="store_true", help="Copy the reviewed candidate into the app. Nothing else.")
    ap.add_argument("--survey", nargs="+", metavar="FILE",
                    help="App survey export(s) to merge into a copy of the sheet. Touches nothing else.")
    args = ap.parse_args()

    if args.survey:
        if args.apply:
            sys.exit("--survey and --apply cannot run together: a person checks the merged sheet first.")
        for path in [args.xlsx] + args.survey:
            if not os.path.exists(path):
                sys.exit("Cannot find: %s" % path)
        merge_survey(args.xlsx, args.sheet, args.survey)
        return

    if args.apply:
        if not os.path.exists(CANDIDATE):
            sys.exit("No candidate yet. Run without --apply first and read the report.")
        shutil.copyfile(CANDIDATE, APP_FILE)
        print("copied  %s\n    ->  %s" % (CANDIDATE, APP_FILE))
        print("The app uses it on the next run. Riders get it on the next update you publish.")
        return

    if not os.path.exists(args.xlsx):
        sys.exit("Cannot find the spreadsheet: %s" % args.xlsx)
    collected_at = args.collected_at or dt.date.fromtimestamp(os.path.getmtime(args.xlsx)).isoformat()

    rows, sheet = read_sheet(args.xlsx, args.sheet)
    shops, facts = build(rows)
    diff = compare_with_app(shops, APP_FILE)

    payload = {
        "collected_at": collected_at,
        "shops": [{k: v for k, v in s.items() if not k.startswith("_")} for s in shops],
    }
    os.makedirs(OUT_DIR, exist_ok=True)
    with open(CANDIDATE, "w", encoding="utf-8") as fh:
        # No spaces: this file rides along on every over-the-air update.
        json.dump(payload, fh, ensure_ascii=False, separators=(",", ":"))
    size = os.path.getsize(CANDIDATE)
    write_report(REPORT, args.xlsx, sheet, collected_at, rows, shops, facts, diff, size)

    print("%d rows -> %d shops, %d skipped, %d duplicate pairs flagged, %d with phones" % (
        len(rows), len(shops), len(facts["skipped"]), len(facts["duplicates"]),
        sum(1 for s in shops if "phones" in s)))
    print("report     %s" % REPORT)
    print("candidate  %s  (%d KB)" % (CANDIDATE, round(size / 1024)))
    print("The app has NOT been changed. Read the report, then run with --apply.")


if __name__ == "__main__":
    main()
